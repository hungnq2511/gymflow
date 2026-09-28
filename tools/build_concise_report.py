from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import build_course_report as b
import build_course_report_7_chapters as old
from report_plain_diagrams import make_all

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'artifacts/Bao_cao_GymFlow_co_dong_Agile_Scrum.docx'
doc = Document()
b.style_document(doc)
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.left_margin, sec.right_margin = Mm(25), Mm(20)
sec.top_margin, sec.bottom_margin = Mm(20), Mm(20)
for r in sec.header.paragraphs[0].runs:
    r.font.color.rgb = b.RGBColor(0, 0, 0)
pages = []
current_major = True

def page(title, major=False):
    global current_major
    current_major=major
    doc.add_page_break()
    doc.add_heading(title, 1 if major else 2)
    pages.append(title)

def p(text):
    return b.add_para(doc, text)

def h(text):
    heading=doc.add_heading(text, 2 if current_major else 3)
    heading.paragraph_format.space_before=Pt(8)
    heading.paragraph_format.space_after=Pt(4)
    for run in heading.runs:
        run.font.size=Pt(12)

def bullets(items):
    b.add_bullets(doc, items)

def table(headers, rows, widths=None):
    widths = widths or ([6.4 / len(headers)] * len(headers))
    result=b.add_table(doc, headers, rows, widths=widths, font_size=10)
    for col,width in zip(result.columns,widths):
        col.width=Inches(width)
    return result

def pic(name, caption, width=6.25):
    b.add_picture(doc, name, width, caption)

# Existing project illustrations are retained with honest captions.
assets = ROOT / 'artifacts/report_7_chapters_assets'
v = make_all(assets)

def simple_flow(filename, title, labels):
    from PIL import Image, ImageDraw
    im = Image.new('RGB', (1800, 920), 'white')
    d = ImageDraw.Draw(im)
    d.text((900, 65), title, anchor='mm', font=old.font(39, True), fill='black')
    for i, label in enumerate(labels):
        y = 140 + i*145
        d.rounded_rectangle((370,y,1430,y+100), radius=15, fill='#EAF2F8', outline='#365F91', width=3)
        d.text((900,y+50), label, anchor='mm', font=old.font(28), fill='black')
        if i < len(labels)-1:
            old.arrow(d, (900,y+102), (900,y+140))
    path=assets/filename
    im.save(path)
    return path

architecture = simple_flow('architecture_plain.png', 'Các phần chính của hệ thống', [
    'Quản lý   •   Nhân viên   •   Hội viên',
    'Giao diện trên trình duyệt', 'Kiểm tra đăng nhập và xử lý yêu cầu',
    'Lưu hồ sơ, gói tập, khoản thu và lượt vào tập',
    'Trả kết quả và thông báo cho người dùng'])
sequence = v['sequence']

# 1–3: front matter.
b.add_cover(doc)
pages.append('Trang bìa')
page('Tóm tắt báo cáo', True)
p('GymFlow là website hỗ trợ quản lý một chi nhánh phòng gym. Hệ thống tập trung hồ sơ hội viên, gói tập, khoản thu và lịch sử vào tập vào cùng một nơi. Quản lý theo dõi hoạt động; nhân viên xử lý công việc tại quầy; hội viên xem thông tin của chính mình.')
p('Đề tài trình bày đủ các bước của môn Công nghệ phần mềm: xác định nhu cầu, phân tích yêu cầu, thiết kế, xây dựng và kiểm thử. Các quy tắc quan trọng gồm không cho dùng gói hết hạn, không trừ lượt hai lần khi quét trùng và giữ lại lịch sử khi hủy khoản thu.')
p('Phần áp dụng Agile Scrum tổ chức công việc thành sáu đợt ngắn, gọi là Sprint. Mỗi đợt có mục tiêu, yêu cầu được chọn, việc phải làm, điều kiện nghiệm thu và phần sản phẩm để trình diễn. Sau mỗi đợt, phản hồi được dùng để điều chỉnh kế hoạch tiếp theo.')
p('Sản phẩm hiện có các nhóm chức năng chính theo phạm vi đề tài. Kết quả kiểm tra được ghi nhận ngày 22 tháng 9 năm 2026 gồm 16 bài kiểm thử tự động đạt, kiểm tra mã nguồn không phát hiện lỗi và tạo bản ứng dụng thành công. Kiểm tra toàn bộ thao tác với cơ sở dữ liệu thật vẫn cần bổ sung trước khi dùng chính thức.')
h('Cách đọc báo cáo')
p('Chương 1 đến chương 5 giải thích bài toán và sản phẩm. Chương 6 trình bày cách tổ chức phát triển theo Scrum bằng các công việc cụ thể của GymFlow. Chương 7 nêu kết quả, hạn chế và hướng hoàn thiện. Các tên tiếng Anh cần thiết được giải thích bằng tiếng Việt ngay trong nội dung.')
page('Mục lục', True)
toc = [('1 Giới thiệu đề tài',4),('2 Khảo sát và phân tích yêu cầu',7),('3 Thiết kế hệ thống',14),('4 Xây dựng hệ thống',22),('5 Kiểm thử hệ thống',28),('6 Áp dụng Agile Scrum',32),('7 Kết luận',48),('Tài liệu tham khảo',50),('Phụ lục A Kịch bản trình diễn',51),('Phụ lục B Đối chiếu yêu cầu và sản phẩm',52)]
for label, n in toc:
    b.add_toc_line(doc,label,n,True)
h('Các nội dung chính trong chương 6')
for label, n in [('6.1 Tổ chức công việc',32),('6.2 Danh sách yêu cầu ưu tiên',33),('6.3 Điều kiện nghiệm thu',35),('6.4 Chia Sprint',36),('6.5 đến 6.10 Kế hoạch và kết quả sáu Sprint',37),('6.11 Theo dõi công việc hằng ngày',43),('6.12 Xử lý thay đổi',44),('6.13 Đánh giá sản phẩm cuối Sprint',45),('6.14 Rút kinh nghiệm',46),('6.15 Tổng hợp kết quả',47)]:
    b.add_toc_line(doc,label,n)

# 4–6: introduction.
page('1 Giới thiệu đề tài', True)
h('1.1 Lý do chọn đề tài')
p('Phòng gym thường tiếp nhận hội viên mới, bán gói, thu tiền và kiểm tra người vào tập mỗi ngày. Nếu dùng sổ hoặc nhiều bảng tính riêng, nhân viên phải tra cứu nhiều lần. Thông tin dễ bị thiếu hoặc không thống nhất, nhất là khi đổi ca làm việc.')
p('Ví dụ, một hội viên đã gia hạn nhưng danh sách tại quầy chưa cập nhật có thể bị từ chối vào tập. Ngược lại, gói đã hết lượt nhưng chưa được ghi nhận kịp thời có thể tiếp tục được sử dụng. Việc sửa một khoản thu mà không giữ lịch sử cũng gây khó khăn khi đối chiếu.')
p('Xây dựng website quản lý giúp các bộ phận sử dụng chung dữ liệu. Nhân viên có thể tìm hội viên, kiểm tra gói và ghi nhận lượt vào ngay tại quầy. Quản lý có thông tin thu chi và hoạt động để theo dõi phòng tập.')
h('Ý nghĩa đối với môn Công nghệ phần mềm')
p('Đề tài có đủ yếu tố để vận dụng kiến thức của môn học: nhiều nhóm người dùng, yêu cầu rõ ràng, dữ liệu liên quan với nhau và những tình huống sai cần kiểm tra. Kết quả được đánh giá qua cả sản phẩm chạy được và cách tổ chức quá trình phát triển.')
p('Agile Scrum phù hợp với việc hoàn thiện từng phần. Sau khi có chức năng quản lý hội viên, nhóm có thể trình diễn và nhận góp ý trước khi xây dựng bán gói. Những yêu cầu chưa rõ được làm rõ sớm, giúp giảm việc sửa lại ở cuối kỳ.')
page('1.2 Mục tiêu và phạm vi')
p('Mục tiêu là xây dựng một website đáp ứng công việc chính của một chi nhánh phòng gym, đồng thời có tài liệu phân tích, thiết kế và kiểm thử để giải thích cách sản phẩm được tạo ra.')
table(['Mục tiêu cụ thể','Kết quả cần có'],[
('Quản lý tập trung','Tra cứu được hồ sơ hội viên và gói đang sử dụng.'),
('Hỗ trợ quầy lễ tân','Bán gói, gia hạn, thu tiền và ghi nhận vào tập.'),
('Kiểm soát dữ liệu','Có quyền sử dụng rõ ràng và lịch sử thao tác.'),
('Hỗ trợ quản lý','Xem khoản thu, chi phí và số liệu tổng hợp.'),
('Thực hành quy trình','Liên kết yêu cầu, kế hoạch Sprint và kiểm thử.')],[2,4.4])
h('Phạm vi thực hiện')
p('Hệ thống phục vụ một chi nhánh. Nhân viên ghi nhận thanh toán đủ bằng tiền mặt hoặc chuyển khoản; phần mềm chưa tự kiểm tra tiền về tài khoản ngân hàng. Hội viên dùng trình duyệt trên điện thoại hoặc máy tính.')
p('Quản lý chuỗi chi nhánh, máy kiểm soát cửa, thanh toán trực tuyến và ứng dụng điện thoại riêng nằm ngoài phạm vi hiện tại. Việc giới hạn này giúp tập trung hoàn thiện chuỗi công việc từ tạo hội viên đến sử dụng gói tập.')
page('1.3 Người sử dụng và cách phát triển')
table(['Người dùng','Công việc chính'],[
('Quản lý','Quản lý gói tập, nhân viên, báo cáo, chi phí và các thao tác cần phê duyệt.'),
('Nhân viên','Tiếp nhận hội viên, bán gói, ghi nhận vào tập và chăm sóc hội viên.'),
('Hội viên','Xem mã vào tập, gói đã mua, lịch sử thanh toán và lịch sử tập.')],[1.4,5])
h('1.4 Công nghệ sử dụng')
p('Next.js và React được dùng để xây dựng website. TypeScript hỗ trợ phát hiện lỗi khi viết chương trình. Supabase cung cấp phần đăng nhập và kết nối dữ liệu; PostgreSQL lưu hồ sơ, gói tập và các giao dịch. Công cụ Vitest dùng để kiểm tra tự động một số quy tắc.')
h('1.5 Phương pháp phát triển')
p('Scrum tổ chức việc phát triển theo các Sprint có độ dài cố định. Trong đề tài, phương án chia việc gồm sáu Sprint, mỗi Sprint dự kiến hai tuần. Nhóm chọn những yêu cầu quan trọng, thực hiện, kiểm tra rồi trình diễn phần sản phẩm đã làm được.')
p('Mỗi chức năng phải có điều kiện nghiệm thu trước khi bắt đầu. Chẳng hạn, chức năng vào tập cần xác định rõ khi nào được chấp nhận, khi nào bị từ chối và có trừ lượt hay không. Cách làm này giúp người viết chương trình và người kiểm tra hiểu cùng một yêu cầu.')

# 7–13: requirements.
page('2 Khảo sát và phân tích yêu cầu', True)
h('2.1 Mô tả bài toán')
p('Chuỗi công việc chính bắt đầu khi khách đăng ký tại quầy. Nhân viên nhập hồ sơ, chọn gói phù hợp, ghi nhận khoản thu và cấp mã hội viên. Khi khách đến tập, nhân viên quét mã để hệ thống kiểm tra quyền sử dụng. Khi gói sắp hết hạn hoặc hết lượt, nhân viên liên hệ để hỗ trợ gia hạn.')
table(['Tình huống','Nhu cầu xử lý'],[
('Khách mới đăng ký','Tạo một hồ sơ và mã nhận diện riêng.'),
('Khách mua hoặc gia hạn','Lưu gói, thời gian sử dụng và khoản thu.'),
('Khách đến tập','Kiểm tra điều kiện và ghi lịch sử vào tập.'),
('Khách tạm nghỉ','Tạm dừng gói theo quyền của quản lý.'),
('Khoản thu bị ghi sai','Hủy có lý do và giữ dấu vết để đối chiếu.'),
('Cuối kỳ cần tổng hợp','Lọc dữ liệu theo thời gian và xem thu chi.')],[2,4.4])
p('Phân tích tập trung vào quy trình được mô hình hóa trong dự án. Các tình huống ở trên là cơ sở xác định yêu cầu; việc đưa vào một phòng gym cụ thể cần xác nhận thêm với người vận hành về quy định bán gói, hoàn tiền và tạm nghỉ.')
page('2.2 Các đối tượng và quyền sử dụng')
p('Ba nhóm người dùng có nhu cầu khác nhau. Quyền được kiểm tra khi thực hiện thao tác, kể cả khi người dùng nhập trực tiếp địa chỉ trang. Việc chỉ ẩn một nút trên màn hình chưa đủ để bảo vệ dữ liệu.')
table(['Chức năng','Quản lý','Nhân viên','Hội viên'],[
('Hồ sơ hội viên','Quản lý','Quản lý','Xem của mình'),('Gói tập','Thêm sửa','Xem để bán','Xem gói đã mua'),
('Bán và gia hạn','Có','Có','Không'),('Ghi nhận vào tập','Có','Có','Xem lịch sử'),
('Tạm dừng gói','Có','Không','Không'),('Hủy khoản thu','Có','Không','Không'),
('Nhân viên và chi phí','Có','Không','Không'),('Báo cáo quản lý','Có','Không','Không'),
('Khu vực cá nhân','Theo vai trò','Theo vai trò','Có')],[2.5,1.2,1.25,1.45])
p('Hội viên chỉ xem dữ liệu liên kết với tài khoản của mình. Nhân viên được xử lý công việc tại quầy nhưng không được tự hủy khoản thu. Quản lý chịu trách nhiệm những thay đổi ảnh hưởng đến tài chính và quyền của người khác.')
page('2.3 Yêu cầu chức năng')
table(['Mã','Yêu cầu','Kết quả người dùng nhận được'],[
('YC01','Đăng nhập và quyền sử dụng','Vào đúng khu vực của từng vai trò.'),
('YC02','Quản lý nhân viên','Tạo tài khoản và quản lý trạng thái.'),
('YC03','Quản lý hội viên','Thêm, sửa, tìm và xem hồ sơ.'),
('YC04','Quản lý gói tập','Khai báo giá, thời hạn và số lượt.'),
('YC05','Bán gói','Lưu quyền sử dụng và khoản thu.'),
('YC06','Gia hạn','Tạo gói kế tiếp theo quy tắc thời gian.'),
('YC07','Ghi nhận vào tập','Cho phép đúng điều kiện, chống quét trùng.'),
('YC08','Tạm dừng gói','Tạm ngừng và mở lại theo quyền quản lý.'),
('YC09','Hủy khoản thu','Lưu lý do hủy và giữ lịch sử.'),
('YC10','Báo cáo','Tổng hợp dữ liệu hoạt động và khoản thu.'),
('YC11','Cổng hội viên','Xem gói, mã vào tập và lịch sử cá nhân.'),
('YC12','Chăm sóc hội viên','Theo dõi gói sắp hết hạn hoặc hết lượt.'),
('YC13','Quản lý chi phí','Ghi chi phí và xem chênh lệch thu chi.')],[.65,2,3.75])
p('Các mã YC được dùng lại trong chương Scrum và phụ lục để đối chiếu. Khi một yêu cầu thay đổi, nhóm cập nhật mô tả và điều kiện nghiệm thu của chính mã đó, tránh tạo nhiều phiên bản không thống nhất.')
page('2.4 Yêu cầu về chất lượng')
table(['Yêu cầu','Cách hiểu trong GymFlow','Cách kiểm tra'],[
('Bảo mật','Người dùng chỉ xem và sửa dữ liệu được phép.','Thử từng vai trò và truy cập hồ sơ người khác.'),
('Dữ liệu chính xác','Bán gói phải lưu đủ gói và khoản thu.','Đối chiếu dữ liệu trước và sau thao tác.'),
('Không ghi trùng','Quét lại không làm trừ lượt nhiều lần.','Thử quét liên tiếp và gửi hai yêu cầu cùng lúc.'),
('Dễ sử dụng','Nhãn tiếng Việt, thông báo nêu rõ cách xử lý.','Thử trên máy tính và điện thoại.'),
('Có thể đối chiếu','Lưu người thực hiện và lý do hủy.','Mở lịch sử sau khi thay đổi.'),
('Dễ bảo trì','Các phần chức năng có cách tổ chức rõ.','Đọc mã nguồn và chạy lại bộ kiểm tra.')],[1.15,2.65,2.6])
p('Tốc độ phản hồi cần phù hợp công việc tại quầy. Báo cáo chưa đưa ra số người dùng đồng thời hoặc thời gian phản hồi đã đạt vì chưa có phép đo tải đầy đủ. Khi thử nghiệm tại phòng gym, nhóm cần đo trên thiết bị và đường truyền thực tế.')
p('Một chức năng chỉ được nghiệm thu khi vừa đúng kết quả, vừa đúng quyền. Ví dụ, hủy khoản thu đúng số tiền nhưng cho phép nhân viên tự thực hiện vẫn là lỗi cần sửa.')
page('2.5 Sơ đồ chức năng Use Case')
pic(v['usecase'],'Hình 2.1 Các nhóm người dùng và chức năng chính')
p('Use Case là cách mô tả một việc người dùng muốn hoàn thành. Trong sơ đồ, hình người là nhóm sử dụng; hình bầu dục là chức năng; đường nối cho biết nhóm nào sử dụng chức năng nào.')
p('Quản lý dùng các chức năng điều hành và kiểm soát. Nhân viên tập trung vào tiếp nhận hội viên, bán gói và ghi nhận vào tập. Hội viên truy cập khu vực cá nhân. Sơ đồ chỉ nêu nhóm chức năng chính; quyền chi tiết thực hiện theo bảng ở mục 2.2.')
h('Cách đọc một tình huống')
p('Với chức năng bán gói, người thực hiện là nhân viên hoặc quản lý. Điều kiện ban đầu là hội viên đã có hồ sơ và gói còn được bán. Kết quả cần đạt là có đăng ký gói, khoản thu và thông tin để đối chiếu.')
page('2.6 Đặc tả đăng nhập và bán gói')
h('Đăng nhập')
p('Người dùng nhập thông tin tài khoản. Hệ thống kiểm tra, xác định vai trò rồi mở đúng khu vực. Nếu thông tin sai hoặc tài khoản đã bị khóa, hệ thống từ chối và thông báo. Nhân viên dùng mật khẩu tạm phải đổi mật khẩu trước khi tiếp tục.')
h('Bán mới và gia hạn gói')
b.add_numbered(doc,[
('Nhân viên tìm đúng hội viên và kiểm tra trạng thái hồ sơ.'),
('Chọn gói còn kinh doanh, ngày bắt đầu và phương thức thanh toán.'),
('Hệ thống kiểm tra quyền, dữ liệu và thời gian sử dụng của gói cũ.'),
('Lưu đăng ký gói, thông tin giá tại thời điểm bán và khoản thu.'),
('Hiển thị kết quả để nhân viên đối chiếu và hướng dẫn hội viên.')])
table(['Tình huống khác','Cách xử lý'],[
('Gói cũ còn hiệu lực','Xác định ngày bắt đầu của gói gia hạn để không chồng thời gian.'),
('Hội viên bị ngừng hoạt động','Từ chối bán theo quy tắc hệ thống.'),
('Thông tin gửi lên không hợp lệ','Báo lỗi và không lưu một phần giao dịch.'),
('Gói thay đổi giá sau khi bán','Giữ giá đã bán trong lịch sử của hội viên.')],[2,4.4])
p('Việc lưu gói và khoản thu phải cùng thành công. Nếu có lỗi giữa quá trình, hệ thống không được để lại gói sử dụng nhưng thiếu khoản thu tương ứng.')
page('2.7 Đặc tả vào tập và hủy khoản thu')
h('Ghi nhận vào tập bằng mã QR hoặc mã hội viên')
p('Nhân viên quét hoặc nhập mã. Hệ thống tìm hội viên, kiểm tra trạng thái, ngày hiệu lực, tình trạng tạm dừng và số lượt còn lại. Nếu đủ điều kiện, hệ thống lưu một lần vào tập và giảm một lượt đối với gói có giới hạn lượt.')
table(['Điều kiện','Phản hồi cần có'],[
('Hội viên đang bị ngừng hoạt động','Từ chối và nêu trạng thái hồ sơ.'),
('Gói chưa bắt đầu hoặc đã hết hạn','Từ chối và hướng dẫn kiểm tra gói.'),
('Gói đang tạm dừng hoặc hết lượt','Từ chối, không trừ thêm lượt.'),
('Quét lại trong thời gian chống trùng','Báo đã ghi nhận, không ghi lần thứ hai.'),
('Gói không giới hạn lượt','Lưu lịch sử nhưng không giảm số lượt.')],[3,3.4])
h('Hủy khoản thu')
p('Quản lý chọn khoản thu còn hiệu lực, nhập lý do và xác nhận. Hệ thống lưu người hủy, thời điểm và lý do, đồng thời xử lý gói liên quan theo quy tắc. Khoản thu đã hủy không được tính vào tổng thu nhưng vẫn xem lại được.')
p('Nhân viên không có quyền hủy. Nếu khoản thu đã bị hủy trước đó, hệ thống thông báo trạng thái hiện tại, tránh thực hiện lần nữa. Cần kiểm tra cả thay đổi khoản thu và gói liên quan khi nghiệm thu tình huống này.')

# 14–21: system design.
page('3 Thiết kế hệ thống', True)
h('3.1 Kiến trúc hệ thống')
pic(architecture,'Hình 3.1 Các phần chính và hướng xử lý yêu cầu',6.15)
p('Hệ thống gồm giao diện, phần xử lý và nơi lưu dữ liệu. Giao diện tiếp nhận thao tác; phần xử lý kiểm tra quyền và quy tắc; nơi lưu dữ liệu giữ hồ sơ cùng lịch sử. Kết quả sau đó được trả về màn hình.')
p('Tách trách nhiệm giúp thay đổi giao diện mà không phải viết lại toàn bộ quy tắc. Những thao tác liên quan đến tiền và lượt tập cần được kiểm tra tại máy chủ để tránh bỏ qua quy định chỉ bằng cách sửa dữ liệu trên trình duyệt.')
page('3.2 Thiết kế các phần chức năng')
table(['Phần','Trách nhiệm','Ví dụ'],[
('Đăng nhập','Xác định người dùng và quyền.','Quản lý được xem báo cáo.'),
('Hội viên và gói tập','Lưu thông tin nền cho nghiệp vụ.','Tìm hội viên theo mã.'),
('Bán gói và khoản thu','Xử lý quyền sử dụng và số tiền.','Tạo gói gia hạn kế tiếp.'),
('Vào tập','Kiểm tra điều kiện và ghi lịch sử.','Từ chối gói đang tạm dừng.'),
('Theo dõi hoạt động','Tổng hợp và hỗ trợ chăm sóc.','Liệt kê gói còn ít lượt.'),
('Khu vực hội viên','Cung cấp thông tin cá nhân.','Xem lịch sử tập của mình.')],[1.7,2.55,2.15])
h('Nguyên tắc xử lý chung')
p('Mỗi thao tác thay đổi dữ liệu đều đi qua bước kiểm tra đăng nhập, quyền và nội dung nhập. Khi không hợp lệ, thông báo cần chỉ rõ vấn đề để người dùng sửa. Khi thành công, màn hình phải cập nhật theo dữ liệu đã lưu.')
p('Các thao tác bán gói, hủy khoản thu và ghi nhận vào tập được gom thành một lần xử lý nhất quán. Cách này hạn chế tình trạng một phần dữ liệu đã thay đổi nhưng phần còn lại chưa thay đổi do mất kết nối hoặc có lỗi.')
h('Tổ chức mã nguồn')
p('Thư mục app chứa các trang và nơi tiếp nhận yêu cầu. components chứa các phần giao diện. lib chứa quy tắc dùng chung. supabase chứa cấu trúc dữ liệu, quy định truy cập và dữ liệu mẫu. Cách chia này hỗ trợ tìm đúng phần cần sửa khi có lỗi.')
page('3.3 Thiết kế cơ sở dữ liệu')
pic(v['erd'],'Hình 3.2 Quan hệ giữa các nhóm dữ liệu chính')
p('Cơ sở dữ liệu là nơi lưu thông tin có cấu trúc. Một hội viên có thể mua nhiều gói theo thời gian. Mỗi đăng ký gói liên kết với loại gói đã mua; khoản thu và lịch sử vào tập liên kết với đăng ký liên quan.')
p('Đường nối thể hiện quan hệ giữa dữ liệu. Ví dụ, một loại gói có thể xuất hiện trong nhiều đăng ký của các hội viên. Một đăng ký gói có nhiều lần vào tập. Mỗi nhóm có mã riêng để liên kết thông tin mà không phải nhập lặp lại toàn bộ hồ sơ.')
p('Hồ sơ hội viên được tách khỏi tài khoản đăng nhập. Nhờ vậy, nhân viên có thể tạo hồ sơ cho khách tại quầy trước, rồi mời khách sử dụng cổng hội viên sau.')
page('3.4 Các dữ liệu và quy tắc lưu trữ')
table(['Nhóm dữ liệu','Thông tin chính','Quy tắc'],[
('Tài khoản','Họ tên, vai trò, trạng thái.','Tài khoản bị khóa không được thao tác.'),
('Hội viên','Mã, liên hệ, mã QR.','Mã hội viên và mã QR không trùng.'),
('Gói tập','Tên, giá, số ngày, số lượt.','Giá không âm, thời hạn hợp lệ.'),
('Đăng ký gói','Hội viên, gói, thời gian, lượt còn.','Giữ thông tin tại lúc mua.'),
('Khoản thu','Số tiền, cách trả, trạng thái.','Giữ lịch sử kể cả khi hủy.'),
('Vào tập','Hội viên, gói, thời điểm.','Không ghi trùng theo quy tắc.'),
('Tạm dừng','Gói, thời gian, trạng thái.','Theo dõi khi mở lại gói.'),
('Nhật ký thao tác','Người làm, nội dung, thời điểm.','Dùng để đối chiếu thay đổi.')],[1.3,2.6,2.5])
p('Giá gói trong danh mục có thể đổi, nhưng khoản thu cũ phải giữ nguyên. Vì vậy, đăng ký gói lưu lại những thông tin quan trọng tại thời điểm bán. Đây là cơ sở giải thích vì sao cùng một loại gói có thể được bán ở các mức giá khác nhau.')
p('Xóa dữ liệu tài chính sẽ làm mất khả năng đối chiếu. Hệ thống dùng trạng thái đã hủy thay cho xóa khoản thu. Báo cáo chỉ cộng các khoản còn hiệu lực.')
page('3.5 Sơ đồ lớp Class Diagram')
pic(v['class'],'Hình 3.3 Các đối tượng chính của bài toán')
p('Sơ đồ lớp mô tả những đối tượng hệ thống cần quản lý. Mỗi ô nêu tên đối tượng, thông tin cần lưu và thao tác liên quan. Sơ đồ thể hiện mô hình của bài toán, không có nghĩa mọi ô đều là một lớp được viết nguyên dạng trong chương trình.')
table(['Đối tượng','Ý nghĩa'],[('Tài khoản và hội viên','Phân biệt thông tin đăng nhập với hồ sơ khách hàng.'),('Loại gói và đăng ký gói','Phân biệt sản phẩm bán ra với gói một khách đã mua.'),('Khoản thu và lần vào tập','Ghi lại giao dịch và quá trình sử dụng.')],[2.8,3.6])
p('Ký hiệu 1 là một đối tượng; 0..1 nghĩa là có thể chưa liên kết hoặc liên kết với một đối tượng. Một đăng ký có thể có lịch sử khoản thu, trong đó hệ thống giới hạn khoản thu còn hiệu lực theo quy tắc. Tên đối tượng trong hình được diễn đạt bằng tiếng Việt để dễ đọc.')
page('3.6 Sơ đồ trình tự vào tập')
pic(sequence,'Hình 3.4 Trình tự trao đổi khi ghi nhận vào tập',6.15)
p('Sơ đồ trình tự, hay Sequence Diagram, cho biết các phần của hệ thống trao đổi theo thứ tự nào. Đọc các mũi tên từ trên xuống để theo dõi từ lúc quét mã đến lúc nhận kết quả. Các đường dọc biểu diễn từng bên tham gia.')
p('Bước quan trọng nằm ở kiểm tra gói và lưu lượt. Hệ thống cần đọc trạng thái mới nhất trước khi cập nhật. Nếu hai yêu cầu đến gần nhau, chỉ yêu cầu hợp lệ mới được ghi, tránh giảm hai lượt cho cùng một lần vào.')
p('Nếu bị từ chối, kết quả trả về phải nêu lý do như hết hạn, hết lượt hoặc đã quét gần đây. Nhân viên nhìn thấy lý do để giải thích cho hội viên; dữ liệu lượt tập không thay đổi khi yêu cầu bị từ chối.')
page('3.7 Sơ đồ hoạt động bán gói')
pic(v['activity'],'Hình 3.5 Các bước bán mới và gia hạn gói')
p('Sơ đồ hoạt động, hay Activity Diagram, mô tả thứ tự công việc và nhánh xử lý. Sau khi chọn hội viên và gói, hệ thống kiểm tra dữ liệu. Dữ liệu hợp lệ mới được lưu; dữ liệu sai dẫn đến thông báo lỗi.')
p('Hình thoi là điểm rẽ nhánh. Nếu quyền hoặc dữ liệu không hợp lệ, hệ thống báo lỗi và không lưu. Nếu hợp lệ, đăng ký gói, khoản thu và nhật ký được ghi cùng nhau. Nếu có lỗi trong lúc lưu, toàn bộ thay đổi của lần bán phải được hoàn lại.')
h('Điều kiện cần bảo đảm')
bullets(['Thông tin hội viên và gói phải tồn tại, còn được phép sử dụng.', 'Ngày bắt đầu và kết thúc tuân theo quy tắc bán mới hoặc gia hạn.', 'Đăng ký gói và khoản thu được lưu cùng nhau.', 'Nhân viên nhận kết quả rõ ràng để đối chiếu sau khi bán.'])
page('3.8 Thiết kế giao diện và bảo vệ dữ liệu')
table(['Màn hình','Cách bố trí'],[
('Tổng quan','Chỉ số chính ở đầu, hoạt động gần đây phía dưới.'),
('Hội viên','Ô tìm kiếm, bộ lọc, danh sách và nút thêm.'),
('Bán gói','Chọn hội viên, gói, ngày bắt đầu và cách thanh toán.'),
('Vào tập','Ưu tiên vùng quét mã và kết quả dễ nhìn.'),
('Cổng hội viên','Mã QR, gói hiện tại và lịch sử cá nhân.')],[1.5,4.9])
p('Các nút, màu trạng thái và cách ghi ngày tháng cần nhất quán. Khi đang lưu, giao diện phải cho người dùng biết hệ thống đang xử lý. Khi không có dữ liệu, cần có hướng dẫn thay vì một vùng trống khó hiểu.')
h('Bảo vệ dữ liệu')
p('Quyền được kiểm tra ở phần xử lý và tại nơi lưu dữ liệu. Khóa quản trị chỉ được dùng trên máy chủ. Mã QR chứa mã ngẫu nhiên để tìm hội viên, không chứa số điện thoại hoặc thông tin cá nhân.')
p('Thao tác ảnh hưởng đến tiền cần xác nhận và ghi lý do. Người quản lý cuối cùng không được vô hiệu hóa theo cách làm hệ thống mất toàn bộ quyền quản trị. Đây là các quy tắc cần thử riêng khi nghiệm thu.')

# 22–27: implementation.
page('4 Xây dựng hệ thống', True)
h('4.1 Công nghệ và vai trò')
table(['Công nghệ','Dùng để làm gì'],[
('Next.js và React','Xây dựng các trang và giao diện website.'),('TypeScript','Kiểm tra cách dùng dữ liệu trong chương trình.'),
('Supabase','Cung cấp đăng nhập và kết nối cơ sở dữ liệu.'),('PostgreSQL','Lưu và xử lý dữ liệu có quan hệ.'),
('Zod','Kiểm tra dữ liệu người dùng gửi lên.'),('Tailwind CSS','Định dạng và điều chỉnh giao diện theo màn hình.'),
('Vitest','Chạy các bài kiểm thử tự động.')],[1.7,4.7])
p('Các công nghệ được kết hợp trong cùng dự án web để thuận tiện phát triển và kiểm tra. Người dùng truy cập bằng trình duyệt; dữ liệu dùng chung được lưu ở cơ sở dữ liệu của hệ thống.')
h('Cách hiện thực một chức năng')
p('Nhóm bắt đầu từ yêu cầu và điều kiện nghiệm thu, thiết kế màn hình, viết phần xử lý, sau đó kiểm tra kết quả. Với bán gói, cần hoàn thành cả màn hình nhập, kiểm tra quyền, lưu gói và lưu khoản thu trước khi xem là xong chức năng.')
p('Tên công nghệ chỉ giải thích công cụ sử dụng. Chất lượng sản phẩm được đánh giá bằng quy tắc hoạt động, dữ liệu đúng và trải nghiệm của người dùng.')
page('4.2 Các chức năng quản lý chính')
h('Hội viên và gói tập')
p('Nhân viên có thể thêm hồ sơ, tìm kiếm, lọc danh sách và cập nhật thông tin. Hệ thống cấp mã hội viên và mã QR để dùng khi vào tập. Dữ liệu trùng hoặc thiếu cần được phát hiện trước khi lưu.')
p('Quản lý tạo gói tập với giá, thời hạn và số lượt nếu có giới hạn. Gói ngừng kinh doanh vẫn được giữ để tra cứu các lần bán trước. Nhân viên xem danh mục đang bán để tư vấn và tạo đăng ký.')
h('Bán gói và thanh toán')
p('Luồng bán lưu thông tin gói đã mua, khoản thu và phiếu thu. Gia hạn xử lý thời gian kế tiếp. Quản lý có thể tạm dừng gói hoặc hủy khoản thu khi đáp ứng điều kiện; mọi thay đổi quan trọng cần giữ lịch sử.')
h('Vào tập và tài khoản cá nhân')
p('Nhân viên dùng camera quét QR hoặc nhập mã. Hệ thống kiểm tra điều kiện trước khi ghi nhận. Hội viên đăng nhập khu vực riêng để xem mã QR, gói hiện tại, gói kế tiếp và lịch sử của mình.')
h('Theo dõi hoạt động')
p('Quản lý xem các chỉ số, khoản thu, chi phí và chênh lệch thu chi. Nhân viên theo dõi danh sách hội viên sắp hết gói, phân công chăm sóc và ghi nội dung liên hệ. Các chức năng này hỗ trợ công việc hằng ngày sau khi luồng bán và vào tập đã có.')
page('4.3 Giao diện tổng quan')
pic(assets/'ui_dashboard.png','Hình 4.1 Minh họa bố cục màn hình tổng quan')
p('Màn hình tổng quan giúp quản lý nhận biết nhanh tình hình phòng tập. Những chỉ số thường xem được đặt trước; danh sách hoạt động nằm bên dưới. Cách bố trí này giảm thao tác chuyển trang khi cần xem một số thông tin chung.')
p('Hình là bản minh họa dựng theo cấu trúc giao diện của dự án, dùng số liệu mẫu. Khi chạy ứng dụng, số liệu được lấy từ dữ liệu đã ghi nhận và quyền của tài khoản đang đăng nhập.')
h('Điểm cần kiểm tra')
bullets(['Khoản thu bị hủy không được cộng vào tổng thu.', 'Khoảng thời gian trên màn hình phải rõ ràng.', 'Tài khoản nhân viên không được mở phần báo cáo dành riêng cho quản lý.', 'Khi chưa có dữ liệu, các chỉ số không gây hiểu nhầm là lỗi tải.'])
page('4.4 Giao diện quản lý hội viên')
pic(assets/'ui_members.png','Hình 4.2 Minh họa bố cục danh sách hội viên')
p('Danh sách tập trung các thông tin cần tra cứu tại quầy như mã hội viên, họ tên, điện thoại và trạng thái. Tìm kiếm và lọc hỗ trợ xác định đúng người trước khi bán gói hoặc xem lịch sử.')
h('Thao tác điển hình')
b.add_numbered(doc,['Nhập từ khóa và chọn đúng hồ sơ.', 'Kiểm tra thông tin liên hệ và trạng thái.', 'Mở chi tiết để xem gói hoặc cập nhật thông tin.', 'Lưu thay đổi và kiểm tra lại trên danh sách.'])
p('Khi thêm hội viên, thông báo dữ liệu không hợp lệ cần đặt gần nội dung nhập. Nếu phát hiện liên hệ trùng, nhân viên nên kiểm tra hồ sơ đã có trước khi tạo mới. Hình dùng dữ liệu mẫu để minh họa bố cục.')
page('4.5 Giao diện ghi nhận vào tập')
pic(assets/'ui_checkin.png','Hình 4.3 Minh họa màn hình ghi nhận vào tập')
p('Màn hình vào tập ưu tiên thao tác nhanh. Nhân viên có thể dùng camera hoặc nhập mã nếu thiết bị không quét được. Kết quả phải cho biết hội viên nào vừa được xử lý và trạng thái có được vào tập hay không.')
table(['Phản hồi','Thông tin cần thể hiện'],[('Được vào tập','Tên hội viên và kết quả ghi nhận.'),('Hết hạn hoặc hết lượt','Lý do từ chối để kiểm tra hoặc gia hạn.'),('Đang tạm dừng','Thông báo tình trạng gói.'),('Quét trùng','Cho biết lượt gần đây đã được ghi nhận.')],[2,4.4])
p('Màu sắc chỉ hỗ trợ nhận biết; thông báo chữ vẫn phải rõ. Hình trên là minh họa giao diện. Kết quả cập nhật số lượt và lịch sử cần được kiểm tra bằng dữ liệu thật trong buổi thử nghiệm.')
page('4.6 Giao diện dành cho hội viên')
pic(assets/'ui_portal.png','Hình 4.4 Minh họa khu vực cá nhân của hội viên')
p('Hội viên xem mã QR để xuất trình tại quầy, gói đang dùng, thời gian còn hiệu lực và lịch sử cá nhân. Việc tự tra cứu giảm nhu cầu hỏi nhân viên về những thông tin đơn giản.')
h('Yêu cầu sử dụng trên điện thoại')
p('Mã QR cần đủ lớn để dễ quét. Thông tin gói cần ngắn gọn và có ngày bắt đầu, ngày kết thúc rõ ràng. Những nội dung dài như lịch sử có thể xem theo danh sách để phù hợp màn hình nhỏ.')
h('Giới hạn truy cập')
p('Mỗi tài khoản chỉ được xem hồ sơ đã liên kết với mình. Cần thử đổi mã hồ sơ trong yêu cầu để xác nhận không đọc được dữ liệu người khác. Việc nhìn thấy đúng màn hình khi đăng nhập chưa đủ để kết luận quyền truy cập đã an toàn.')
p('Hình sử dụng dữ liệu mẫu và thể hiện bố cục tham khảo của chức năng hiện có trong dự án.')

# 28–31: testing.
page('5 Kiểm thử hệ thống', True)
h('5.1 Phương pháp kiểm thử')
p('Kiểm thử bắt đầu từ yêu cầu. Với mỗi chức năng, nhóm xác định trường hợp hợp lệ, dữ liệu sai và tình huống ở giới hạn. Người kiểm tra ghi bước thực hiện, kết quả mong đợi, kết quả thực tế và lỗi cần sửa.')
table(['Cách kiểm thử','Mục đích','Áp dụng'],[
('Kiểm tra từng quy tắc','Xem một quy tắc có trả đúng kết quả.','Hết hạn, hết lượt, tổng thu chi.'),
('Kiểm tra các phần phối hợp','Xem giao diện, xử lý và dữ liệu có khớp.','Bán gói đi cùng khoản thu.'),
('Kiểm tra quyền','Ngăn thao tác vượt quyền.','Nhân viên không được hủy thu.'),
('Thử toàn bộ thao tác','Kiểm tra từ góc nhìn người dùng.','Tạo hội viên đến vào tập.'),
('Kiểm tra lại sau sửa','Đảm bảo sửa lỗi không làm hỏng việc cũ.','Chạy lại bài kiểm thử liên quan.')],[1.55,2.5,2.35])
p('Bộ kiểm thử tự động hiện có tập trung vào một số quy tắc nhỏ. Các tình huống kết nối cơ sở dữ liệu, thao tác đồng thời và toàn bộ màn hình phải được kiểm tra bổ sung. Hai nhóm kết quả được ghi riêng để tránh hiểu rằng mọi chức năng đều đã được thử đầy đủ.')
page('5.2 Các trường hợp kiểm tra vào tập')
p('Bảng dưới đối chiếu với chín trường hợp tự động trong phần kiểm tra quy tắc vào tập. Kết quả đạt ở mức quy tắc; chưa chứng minh camera, đường truyền và dữ liệu thật cùng hoạt động đúng.')
table(['Mã','Dữ liệu thử','Kết quả mong đợi'],[
('KT01','Hồ sơ hợp lệ, gói hiệu lực, còn 5 lượt.','Cho vào và yêu cầu trừ một lượt.'),
('KT02','Gói không giới hạn lượt.','Cho vào, không trừ lượt.'),
('KT03','Hội viên ngừng hoạt động.','Từ chối.'),('KT04','Không có gói đang hoạt động.','Từ chối.'),
('KT05','Ngày bắt đầu ở tương lai.','Từ chối vì chưa đến ngày.'),('KT06','Gói đang tạm dừng.','Từ chối.'),
('KT07','Gói đã hết hạn.','Từ chối.'),('KT08','Gói còn 0 lượt.','Từ chối vì hết lượt.'),
('KT09','Vừa ghi nhận cách đó 5 phút.','Từ chối quét trùng.')],[.6,3.6,2.2])
h('Tình huống cần thử bổ sung')
p('Cần gửi hai yêu cầu gần như cùng lúc để xem dữ liệu có bị trừ hai lần hay không. Cũng cần thử tại thời điểm chuyển ngày, ngay khi hết thời gian chống trùng và khi mất kết nối. Những trường hợp này chưa được tính vào chín kết quả tự động trên.')
page('5.3 Kiểm tra nghiệp vụ và quyền sử dụng')
table(['Mã','Tình huống','Kết quả cần đạt'],[
('KT10','Tạo hồ sơ có điện thoại trùng.','Không tạo thêm hồ sơ trùng.'),
('KT11','Nhân viên thêm gói hoặc hủy thu.','Từ chối do không đủ quyền.'),
('KT12','Bán gói mới hợp lệ.','Có đăng ký gói và khoản thu phù hợp.'),
('KT13','Gia hạn khi gói cũ còn hạn.','Gói kế tiếp có thời gian phù hợp.'),
('KT14','Quản lý hủy thu có lý do.','Giữ lịch sử và cập nhật dữ liệu liên quan.'),
('KT15','Hội viên đọc hồ sơ người khác.','Không trả dữ liệu của người khác.'),
('KT16','Thu 1 triệu, chi 250 nghìn; có khoản hủy.','Chênh lệch 750 nghìn; bỏ khoản hủy.'),
('KT17','Gói còn 3 lượt hoặc gần hết hạn.','Được nhận diện để chăm sóc.')],[.6,3.1,2.7])
p('KT16 và KT17 có bài kiểm tra tự động cho phần tính toán tương ứng. KT10 đến KT15 là kịch bản cần chạy với hệ thống và cơ sở dữ liệu kiểm thử; báo cáo không ghi các trường hợp này là đã đạt khi chưa có biên bản chạy.')
h('Ghi nhận lỗi')
p('Một lỗi cần có mã, bước lặp lại, kết quả mong đợi, kết quả thực tế và mức ảnh hưởng. Sau khi sửa, người kiểm tra làm lại đúng kịch bản gây lỗi. Lỗi liên quan tiền hoặc quyền truy cập phải được ưu tiên trước các chỉnh sửa hình thức.')
page('5.4 Kết quả kiểm tra và đánh giá')
table(['Nhóm kiểm tra','Kết quả ghi nhận','Ý nghĩa'],[
('Quy tắc vào tập','9 bài đạt.','Bao phủ điều kiện chấp nhận và từ chối cơ bản.'),
('Chăm sóc và thu chi','4 bài đạt.','Kiểm tra nhận diện cần chăm sóc và phép tính thu chi.'),
('Kiểm tra mã dữ liệu','3 bài đạt.','Chấp nhận mã hợp lệ và từ chối mã sai.'),
('Kiểm tra mã nguồn','Không phát hiện lỗi.','Đạt các quy tắc kiểm tra đang cấu hình.'),
('Tạo bản ứng dụng','Thành công.','Chương trình được xử lý thành bản có thể chạy.'),
('Thao tác đầy đủ với dữ liệu thật','Cần bổ sung.','Chưa đủ căn cứ kết luận toàn hệ thống đạt.')],[1.8,1.8,2.8])
p('Tổng cộng có 16 bài kiểm thử tự động đạt trong lần kiểm chứng ngày 22 tháng 9 năm 2026. Kết quả này xác nhận phần đã được kiểm tra, không thay thế việc chạy thử toàn bộ chức năng tại phòng gym.')
h('Điều kiện trước khi vận hành')
bullets(['Chạy kịch bản đầy đủ cho cả ba vai trò.', 'Thử bán gói, gia hạn, tạm dừng và hủy thu với dữ liệu kiểm thử.', 'Xác nhận không đọc được dữ liệu hội viên khác.', 'Thử mất kết nối, thao tác đồng thời, sao lưu và khôi phục.', 'Lưu kết quả thực tế và xử lý lỗi nghiêm trọng còn lại.'])

# 32–47: concrete Scrum application, not invented process history.
page('6 Áp dụng Agile Scrum', True)
h('6.1 Tổ chức công việc của đề tài')
p('Mục tiêu sản phẩm là giúp một chi nhánh quản lý được chuỗi công việc từ tiếp nhận hội viên, bán gói đến ghi nhận vào tập và theo dõi hoạt động. Scrum được áp dụng bằng việc chia mục tiêu này thành các phần có thể kiểm tra và trình diễn sau từng Sprint [1].')
p('Phương án tổ chức gồm sáu Sprint, mỗi Sprint dự kiến hai tuần. Lịch và phân công trong chương là kế hoạch đề xuất cho đề tài; phần kết quả đối chiếu với chức năng hiện có và bằng chứng kiểm thử. Lịch sử họp, số giờ thực tế và việc nghiệm thu từng Sprint cần được xác nhận bằng hồ sơ thực hiện.')
table(['Trách nhiệm trong Scrum','Cách áp dụng cho GymFlow'],[
('Product Owner\nNgười phụ trách yêu cầu','Xếp thứ tự yêu cầu, làm rõ quy tắc bán gói và quyết định giá trị cần ưu tiên.'),
('Scrum Master\nNgười hỗ trợ quy trình','Giúp duy trì nhịp làm việc, làm rõ vướng mắc và thúc đẩy cải tiến.'),
('Developers\nNhóm thực hiện','Cùng phân tích, thiết kế, lập trình, kiểm tra và tạo phần sản phẩm sử dụng được.')],[2.1,4.3])
p('Phân công cụ thể cần gắn với thành viên thực tế. Với bài làm cá nhân, người thực hiện có thể dùng cách lập kế hoạch và tự kiểm tra theo Sprint, nhưng cần trình bày rõ mức độ vận dụng thay vì xem đó là một nhóm Scrum đầy đủ.')
page('6.2 Product Backlog phần yêu cầu cốt lõi')
p('Product Backlog là danh sách yêu cầu được sắp theo mức quan trọng. Mỗi dòng mô tả điều người dùng muốn làm và giá trị nhận được. Ưu tiên được xét theo công việc thiết yếu tại quầy, mức ảnh hưởng nếu sai và quan hệ với chức năng khác.')
table(['Mã','Nhu cầu người dùng','Ưu tiên','Sprint'],[
('YC01','Là người dùng, tôi muốn đăng nhập đúng quyền để sử dụng đúng chức năng.','Cao','1'),
('YC02','Là quản lý, tôi muốn tạo và khóa tài khoản nhân viên để kiểm soát người thao tác.','Cao','1'),
('YC03','Là nhân viên, tôi muốn tìm và cập nhật hội viên để phục vụ đúng người.','Cao','2'),
('YC04','Là quản lý, tôi muốn thiết lập gói tập để nhân viên bán đúng giá và điều kiện.','Cao','2'),
('YC05','Là nhân viên, tôi muốn bán gói và ghi thu để khách có quyền sử dụng.','Cao','3'),
('YC06','Là nhân viên, tôi muốn gia hạn để khách tiếp tục tập đúng thời gian.','Cao','3'),
('YC09','Là quản lý, tôi muốn hủy khoản thu sai có lý do để số liệu còn đối chiếu được.','Cao','3')],[.6,4.2,.8,.8])
p('Đăng nhập và dữ liệu hội viên được làm trước vì các chức năng phía sau đều cần. Bán gói và hủy thu được xem xét cùng nhau để thống nhất cách ghi nhận doanh thu và xử lý sai sót.')
page('6.2 Product Backlog phần vận hành và chất lượng')
table(['Mã','Nhu cầu hoặc công việc','Ưu tiên','Sprint'],[
('YC07','Nhân viên ghi nhận vào tập đúng điều kiện, không trừ lượt trùng.','Cao','4'),
('YC08','Quản lý tạm dừng và mở lại gói theo quy định.','Vừa','4'),
('YC10','Quản lý xem báo cáo để theo dõi hoạt động.','Vừa','5'),
('YC11','Hội viên tự xem gói, QR và lịch sử của mình.','Vừa','5'),
('YC12','Nhân viên nhận diện và liên hệ hội viên cần gia hạn.','Vừa','6'),
('YC13','Quản lý ghi chi phí và đối chiếu thu chi.','Vừa','6'),
('CL01','Kiểm tra quyền và dữ liệu ở mỗi chức năng.','Cao','1–6'),
('CL02','Kiểm tra lại chức năng cũ sau mỗi thay đổi.','Cao','1–6'),
('CL03','Bổ sung thử toàn hệ thống, tải và khôi phục dữ liệu.','Cao','Tiếp theo')],[.6,4.2,.8,.8])
h('Cách quản lý danh sách')
p('Người phụ trách yêu cầu cập nhật nội dung sau khi nhận phản hồi. Nhóm thực hiện đánh giá việc cần làm và khả năng hoàn thành trong Sprint. Các yêu cầu chưa rõ phải được làm rõ trước khi đưa vào kế hoạch chi tiết.')
p('CL01 và CL02 là công việc chất lượng đi cùng từng chức năng, không chờ đến Sprint cuối mới bắt đầu. CL03 là phần còn cần bổ sung theo kết quả hiện có; cần ưu tiên trước khi đưa sản phẩm vào vận hành chính thức.')
page('6.3 Tiêu chí nghiệm thu và điều kiện hoàn thành')
p('Tiêu chí nghiệm thu nêu kết quả quan sát được của từng yêu cầu. Nhờ viết trước, nhóm tránh trường hợp giao diện đã có nhưng quy tắc quan trọng chưa làm. Các ví dụ dưới đây được dùng làm cơ sở kiểm tra cuối Sprint.')
table(['Yêu cầu','Điều kiện nghiệm thu tiêu biểu'],[
('YC03 Hội viên','Tạo hồ sơ hợp lệ, tìm lại được; báo rõ dữ liệu thiếu hoặc trùng.'),
('YC05 Bán gói','Một lần bán tạo đúng gói và khoản thu; lỗi không để lại dữ liệu dở dang.'),
('YC07 Vào tập','Gói hợp lệ được vào; hết hạn, hết lượt hoặc quét trùng bị từ chối.'),
('YC09 Hủy thu','Chỉ quản lý được hủy, bắt buộc lý do và còn lịch sử.'),
('YC11 Cổng hội viên','Xem đúng dữ liệu cá nhân và bị chặn khi yêu cầu dữ liệu người khác.'),
('YC13 Thu chi','Chỉ cộng khoản còn hiệu lực; phép tính khớp dữ liệu mẫu.')],[1.6,4.8])
h('Điều kiện hoàn thành chung')
bullets(['Chức năng đáp ứng tiêu chí đã thống nhất và được kiểm tra quyền.', 'Các phần liên quan đã kết nối; dữ liệu lưu đúng và thông báo rõ.', 'Bài kiểm tra liên quan được chạy; lỗi nghiêm trọng đã được xử lý.', 'Mã được kiểm tra, tạo được bản ứng dụng và có hướng dẫn trình diễn.', 'Bằng chứng kiểm tra được lưu để người khác có thể đối chiếu.'])
p('Đây là thỏa thuận chất lượng của đề tài. Hạng mục chưa đáp ứng không được ghi là hoàn thành chỉ vì đã viết xong mã. Kết quả kiểm tra hiện có được đối chiếu riêng ở mục 6.15.')
page('6.4 Chia Sprint và lập kế hoạch')
table(['Sprint','Thời gian dự kiến','Mục tiêu','Yêu cầu'],[
('1','Tuần 1–2','Đăng nhập và sử dụng đúng quyền.','YC01–02'),
('2','Tuần 3–4','Quản lý được hội viên và danh mục gói.','YC03–04'),
('3','Tuần 5–6','Bán, gia hạn và xử lý khoản thu.','YC05–06,09'),
('4','Tuần 7–8','Kiểm soát sử dụng gói khi vào tập.','YC07–08'),
('5','Tuần 9–10','Xem báo cáo và thông tin cá nhân.','YC10–11'),
('6','Tuần 11–12','Chăm sóc hội viên và theo dõi thu chi.','YC12–13')],[.6,1.25,3.1,1.45])
p('Đây là khung thời gian đề xuất, không phải mốc ngày đã diễn ra. Mỗi Sprint đều có phân tích, thiết kế, xây dựng và kiểm thử cho phần việc được chọn. Nhóm có thể trình diễn một luồng sử dụng cuối đợt.')
h('Cách chọn khối lượng')
p('Đầu Sprint, nhóm xác định mục tiêu và chọn yêu cầu vừa với thời gian có thể làm. Công việc lớn được tách thành nhiệm vụ ngắn có kết quả kiểm tra được. Cần chừa thời gian sửa lỗi và trình diễn, không dùng toàn bộ thời gian cho lập trình.')
p('Kế hoạch từng Sprint ở các trang sau chia thành mười ngày làm việc dự kiến. Các khoảng ngày là hướng sắp xếp, không phải số công lao động đã đo. Sau mỗi đợt, nhóm dùng công việc thực tế hoàn tất để điều chỉnh khối lượng đợt sau.')

sprints = [
dict(n=1,title='Đăng nhập và quyền sử dụng', ids='YC01 và YC02', goal='Người dùng vào đúng khu vực; quản lý có thể kiểm soát tài khoản nhân viên.',
 tasks=[('1–2','Làm rõ quyền của ba vai trò; chuẩn bị tài khoản thử.'),('3–5','Xây dựng đăng nhập, đổi mật khẩu và quản lý nhân viên.'),('6–8','Kiểm tra tài khoản sai, bị khóa và quyền truy cập trực tiếp.'),('9–10','Sửa lỗi, trình diễn luồng đăng nhập và ghi phản hồi.')],
 demo='Đăng nhập lần lượt bằng quản lý, nhân viên và hội viên; thử mở một chức năng không được phép; thử đổi mật khẩu tạm.',
 result='Mã nguồn có các trang đăng nhập, khôi phục và đổi mật khẩu; có kiểm tra vai trò cùng phần quản lý nhân viên. Đây là nền tảng cho các chức năng tiếp theo.',
 remain='Cần lưu kết quả chạy thử các vai trò trên môi trường dữ liệu kiểm thử; xác nhận luồng mời tài khoản và khôi phục mật khẩu bằng email.',
 lesson='Chốt bảng quyền trước khi viết giao diện giúp tránh phải sửa lại nhiều màn hình.'),
dict(n=2,title='Hội viên và danh mục gói', ids='YC03 và YC04', goal='Nhân viên tìm được đúng hội viên; quản lý thiết lập được gói để chuẩn bị bán.',
 tasks=[('1–2','Chốt thông tin bắt buộc và cách xử lý liên hệ trùng.'),('3–5','Làm danh sách, tìm kiếm, thêm sửa hội viên và mã QR.'),('6–8','Làm danh mục gói; thử giá, thời hạn và quyền sửa gói.'),('9–10','Thử dữ liệu sai, sửa lỗi và trình diễn từ hồ sơ đến gói.')],
 demo='Tạo hội viên, tìm lại theo liên hệ, cập nhật hồ sơ; quản lý tạo gói rồi ngừng bán; nhân viên thử thao tác chỉ dành cho quản lý.',
 result='Dự án có danh sách hội viên, thao tác cập nhật, mã nhận diện và danh mục gói tập. Các chức năng này cung cấp dữ liệu cho luồng bán gói.',
 remain='Cần thử liên hệ trùng, thông tin thiếu và việc ngừng bán gói đã từng được sử dụng. Không suy ra các trường hợp này đạt chỉ từ việc có màn hình.',
 lesson='Thông báo lỗi nên nói rõ trường cần sửa để nhân viên không phải nhập lại toàn bộ hồ sơ.'),
dict(n=3,title='Bán gói gia hạn và khoản thu', ids='YC05, YC06 và YC09', goal='Một lần bán tạo đủ gói và khoản thu; các sai sót có thể xử lý mà vẫn giữ lịch sử.',
 tasks=[('1–2','Thống nhất ngày gia hạn, giá lưu tại lúc bán và lý do hủy.'),('3–5','Xây dựng bán mới, gia hạn và lưu khoản thu.'),('6–8','Làm hủy thu; thử quyền, dữ liệu sai và tính nhất quán.'),('9–10','Đối chiếu dữ liệu, sửa lỗi và trình diễn cả nhánh hủy.')],
 demo='Bán gói cho một hội viên; gia hạn khi gói còn hiệu lực; quản lý hủy một khoản thu có lý do và kiểm tra lịch sử liên quan.',
 result='Mã nguồn có xử lý bán gói, gia hạn, tạo khoản thu và hủy khoản thu. Những thay đổi liên quan được gom trong phần xử lý dữ liệu để tránh lưu dở dang.',
 remain='Cần chạy thử trọn luồng trên cơ sở dữ liệu, đặc biệt lúc có lỗi giữa quá trình và khi gửi lặp yêu cầu. Chưa có căn cứ nghiệm thu toàn bộ luồng chỉ bằng kiểm thử quy tắc nhỏ.',
 lesson='Kế hoạch phải dành thời gian cho trường hợp sai và hủy giao dịch vì ảnh hưởng trực tiếp đến số liệu.'),
dict(n=4,title='Vào tập và tạm dừng gói', ids='YC07 và YC08', goal='Chỉ ghi nhận lượt vào hợp lệ và phản hồi rõ lý do từ chối.',
 tasks=[('1–2','Liệt kê điều kiện vào tập và quy tắc tạm dừng.'),('3–5','Làm quét hoặc nhập mã; ghi lịch sử và cập nhật lượt.'),('6–8','Làm tạm dừng, mở lại; viết kiểm tra quy tắc vào tập.'),('9–10','Thử quét trùng, đối chiếu lượt và trình diễn tình huống lỗi.')],
 demo='Cho hội viên có gói hợp lệ vào tập, quét lại ngay, thử gói hết lượt và gói tạm dừng; so sánh lịch sử trước và sau.',
 result='Dự án có luồng vào tập, kiểm soát quét trùng và tạm dừng gói. Chín bài kiểm tra tự động cho quyết định vào tập đã đạt trong lần kiểm chứng được ghi nhận.',
 remain='Cần thử camera, hai yêu cầu đồng thời và cập nhật dữ liệu thật. Chín bài tự động kiểm tra quy tắc, chưa kiểm tra toàn bộ thiết bị và quá trình lưu.',
 lesson='Thiết kế ca kiểm tra trước giúp làm rõ điều kiện từ chối và tránh chỉ thử trường hợp được vào tập.'),
dict(n=5,title='Báo cáo và cổng hội viên', ids='YC10 và YC11', goal='Quản lý theo dõi được số liệu; hội viên tự xem được thông tin cá nhân.',
 tasks=[('1–2','Chốt chỉ số cần xem và thông tin riêng của hội viên.'),('3–5','Làm tổng quan, báo cáo và cách lọc dữ liệu.'),('6–8','Làm cổng hội viên; thử quyền xem dữ liệu và điện thoại.'),('9–10','Đối chiếu số liệu, sửa lỗi và trình diễn hai vai trò.')],
 demo='Quản lý mở báo cáo theo thời gian; hội viên mở mã QR và lịch sử; thử yêu cầu dữ liệu của một hội viên khác.',
 result='Mã nguồn có tổng quan, báo cáo và khu vực cá nhân. Hội viên có thể được cung cấp gói đang dùng, gói kế tiếp và lịch sử theo tài khoản liên kết.',
 remain='Cần lập dữ liệu có số tổng biết trước để đối chiếu báo cáo; chạy thử truy cập chéo giữa hai hội viên và nhiều cỡ màn hình.',
 lesson='Ưu tiên tính đúng và quyền xem trước khi bổ sung nhiều biểu đồ; một chỉ số sai có thể làm người quản lý hiểu sai hoạt động.'),
dict(n=6,title='Chăm sóc hội viên và thu chi', ids='YC12 và YC13', goal='Nhân viên theo dõi khách cần gia hạn; quản lý đối chiếu được khoản thu và chi phí.',
 tasks=[('1–2','Xác định gói cần chăm sóc và loại chi phí cần lưu.'),('3–5','Làm phân công liên hệ, lịch sử chăm sóc và nhập chi phí.'),('6–8','Kiểm tra phép tính thu chi, chạy lại các bài kiểm tra cũ.'),('9–10','Tạo bản ứng dụng, hoàn thiện tài liệu và trình diễn tổng thể.')],
 demo='Tìm hội viên sắp hết gói, ghi nội dung liên hệ; nhập chi phí; kiểm tra tổng thu, tổng chi và chênh lệch với dữ liệu mẫu.',
 result='Dự án có chăm sóc hội viên và chi phí. Bốn bài kiểm tra tự động cho chăm sóc và thu chi đạt; cùng các nhóm khác tạo thành 16 bài đạt. Kiểm tra mã và tạo bản ứng dụng thành công.',
 remain='Còn cần thử đầy đủ các luồng với dữ liệu thật, sao lưu, khôi phục và tải. Đây là công việc tiếp tục trước vận hành, không nên đánh dấu hoàn tất toàn bộ chất lượng.',
 lesson='Đợt cuối vẫn tạo chức năng sử dụng được, đồng thời tổng hợp bằng chứng; tránh dành cả đợt chỉ để viết báo cáo.')]

for s in sprints:
    page(f'6.{s["n"]+4} Sprint {s["n"]} {s["title"]}')
    p(f'Mục tiêu: {s["goal"]} Yêu cầu được chọn: {s["ids"]}.')
    table(['Ngày dự kiến','Công việc và đầu ra'],s['tasks'],[1.05,5.35])
    h('Trình diễn và nghiệm thu')
    p(s['demo'])
    h('Kết quả đối chiếu sản phẩm hiện có')
    p(s['result'])
    h('Việc cần kiểm tra tiếp')
    p(s['remain'])
    p('Điểm cần lưu ý cho cách làm: '+s['lesson'])

page('6.11 Theo dõi công việc hằng ngày')
p('Trong mỗi Sprint, nhóm dùng bảng việc gồm Chưa làm, Đang làm, Chờ kiểm tra và Hoàn thành. Mỗi việc có một người chịu trách nhiệm cập nhật. Người viết chương trình và người kiểm tra phối hợp để tránh dồn nhiều việc chưa được thử vào cuối đợt.')
table(['Thẻ việc minh họa','Yêu cầu','Dấu hiệu hoàn thành'],[
('Tách điều kiện vào tập','YC07','Có danh sách điều kiện và kết quả mong đợi.'),
('Hiển thị lý do từ chối','YC07','Người dùng hiểu được gói hết hạn, hết lượt hoặc quét trùng.'),
('Kiểm tra trừ lượt','YC07','Dữ liệu trước và sau khớp số lượt cần giảm.'),
('Thử quét đồng thời','YC07','Không ghi hai lần cho cùng một lượt hợp lệ.')],[2, .7,3.7])
p('Bảng trên là cách tách việc cho Sprint 4, không phải ảnh chụp tiến độ thực tế. Khi áp dụng, từng thẻ phải có trạng thái, người thực hiện và bằng chứng kiểm tra của nhóm.')
h('Trao đổi ngắn mỗi ngày')
p('Daily Scrum là buổi trao đổi ngắn để nhóm thực hiện xem tiến độ hướng đến mục tiêu Sprint và điều chỉnh kế hoạch [1]. Với GymFlow, nội dung cần tập trung vào chức năng đã kiểm tra được, việc đang vướng và việc cần phối hợp tiếp theo.')
p('Ví dụ, nếu camera chưa hoạt động, nhóm có thể tiếp tục thử phần nhập mã và quy tắc vào tập, đồng thời giao việc kiểm tra camera cho người phụ trách. Vướng mắc phải được ghi rõ và xử lý, không chỉ báo rằng công việc đang chậm.')
page('6.12 Xử lý yêu cầu thay đổi và rủi ro')
p('Yêu cầu mới được ghi vào danh sách chung, làm rõ giá trị và ảnh hưởng trước khi đưa vào Sprint. Nhóm cùng người phụ trách yêu cầu xem xét để giữ mục tiêu đợt đang làm. Lỗi làm sai tiền hoặc lộ dữ liệu cần được xử lý ngay theo mức ảnh hưởng.')
table(['Tình huống minh họa','Cách xử lý','Ảnh hưởng kế hoạch'],[
('Đề nghị thêm quản lý nhiều chi nhánh','Ghi yêu cầu, xác định dữ liệu và quyền mới.','Đưa vào hướng phát triển sau phạm vi một chi nhánh.'),
('Muốn đổi cách gia hạn','Làm rõ ngày bắt đầu và các gói đã bán.','Cập nhật YC06, bổ sung kiểm tra trước khi sửa.'),
('Phát hiện nhân viên hủy được thu','Xem là lỗi quyền cần ưu tiên.','Sửa, kiểm tra lại và điều chỉnh việc ít quan trọng.'),
('Cuối Sprint còn chức năng chưa kiểm tra','Giữ trạng thái chưa hoàn thành.','Đưa phần còn lại về danh sách và lập kế hoạch lại.')],[2,2.5,1.9])
h('Rủi ro cần theo dõi')
p('Các rủi ro chính là thiếu môi trường dữ liệu thử, quy tắc chưa thống nhất và khối lượng quá lớn. Biện pháp là chuẩn bị dữ liệu từ đầu, xác nhận ví dụ nghiệp vụ trước khi viết mã và chia yêu cầu lớn thành phần có thể trình diễn.')
p('Tình huống trong bảng dùng để minh họa cách ra quyết định. Khi có thay đổi thực tế, cần lưu ngày, người đề nghị, quyết định và tiêu chí nghiệm thu được sửa để giải thích vì sao kế hoạch thay đổi.')
page('6.13 Đánh giá sản phẩm cuối Sprint')
p('Cuối Sprint, nhóm trình diễn chức năng theo tình huống người dùng. Người phụ trách yêu cầu và bên liên quan cùng xem kết quả, góp ý và điều chỉnh việc tiếp theo. Buổi này gọi là Sprint Review [1]. Việc đánh giá dùng dữ liệu và thao tác cụ thể.')
table(['Sprint','Nội dung trình diễn','Bằng chứng cần lưu'],[
('1','Đăng nhập đúng vai trò và từ chối vượt quyền.','Tài khoản thử và kết quả từng vai trò.'),
('2','Tạo hội viên, tìm lại và thiết lập gói.','Dữ liệu trước sau, tình huống nhập sai.'),
('3','Bán, gia hạn và hủy thu có lý do.','Đăng ký gói, khoản thu và lịch sử.'),
('4','Vào tập hợp lệ và các nhánh từ chối.','Số lượt, lịch sử và kết quả kiểm tra.'),
('5','Báo cáo quản lý và dữ liệu hội viên.','Số tổng đối chiếu và kiểm tra quyền.'),
('6','Chăm sóc, chi phí và chạy luồng tổng thể.','Dữ liệu liên hệ, phép tính và báo cáo kiểm tra.')],[.65,3,2.75])
h('Cách ghi kết quả đánh giá')
p('Biên bản cần nêu mục tiêu, yêu cầu đã trình diễn, ý kiến nhận được và việc cần sửa. Với mỗi tiêu chí, ghi Đạt, Chưa đạt hoặc Chưa kiểm tra. Hạng mục chỉ mới có mã nguồn phải được phân biệt với hạng mục đã chạy và đáp ứng điều kiện hoàn thành.')
p('Báo cáo hiện ghi nhận kết quả kiểm tra mã và 16 bài tự động; các biên bản trình diễn từng Sprint cần được lập từ buổi làm việc thực tế. Không dùng kết quả một lần chạy cuối kỳ để suy ra rằng mọi Sprint đều đã được nghiệm thu.')
page('6.14 Rút kinh nghiệm và cải tiến cách làm')
p('Sau khi xem sản phẩm, nhóm dành thời gian xem lại cách làm việc để chọn một số cải tiến cho đợt tiếp theo. Buổi rút kinh nghiệm gọi là Sprint Retrospective [1]. Nội dung cần dẫn đến hành động cụ thể, có người theo dõi và cách kiểm tra hiệu quả.')
table(['Vấn đề cần xem xét','Hành động cải tiến đề xuất','Cách theo dõi'],[
('Yêu cầu được hiểu khác nhau','Viết ví dụ hợp lệ và không hợp lệ trước khi làm.','Mỗi yêu cầu được chọn có tiêu chí rõ.'),
('Kiểm tra dồn cuối đợt','Tạo ca kiểm tra khi chia việc, thử ngay khi xong phần nhỏ.','Theo dõi số việc chờ kiểm tra.'),
('Khó tìm nguyên nhân lỗi','Ghi bước gây lỗi và dữ liệu đầu vào.','Người khác lặp lại được lỗi.'),
('Nhiều việc bắt đầu nhưng chưa xong','Ưu tiên hoàn tất việc đang làm trước khi mở thêm.','Đếm việc hoàn thành theo tiêu chí.'),
('Báo cáo tiến độ thiếu căn cứ','Gắn kết quả kiểm tra vào thẻ việc.','Mỗi việc hoàn thành có bằng chứng.')],[1.8,2.8,1.8])
p('Các hành động trên là đề xuất phù hợp với đặc điểm dự án, chưa phải biên bản họp đã diễn ra. Khi dùng thực tế, nhóm chọn ít hành động có thể thực hiện ngay thay vì đặt quá nhiều mục tiêu cải tiến.')
h('Ví dụ áp dụng')
p('Sau Sprint bán gói, nhóm có thể chọn cải tiến là chuẩn bị dữ liệu lỗi trước khi lập trình Sprint vào tập. Kết quả cần thấy là mỗi nhánh từ chối đều có tình huống thử. Có thể đối chiếu với chín bài kiểm tra quy tắc vào tập hiện có để đánh giá mức bao phủ.')
page('6.15 Tổng hợp kết quả và giá trị áp dụng')
table(['Nhóm việc','Kết quả sản phẩm','Mức bằng chứng'],[
('Sprint 1–2','Đăng nhập, nhân viên, hội viên và gói tập.','Có mã nguồn; cần hồ sơ nghiệm thu đầy đủ.'),
('Sprint 3','Bán gói, gia hạn, hủy khoản thu.','Có xử lý; cần thử trọn luồng dữ liệu.'),
('Sprint 4','Vào tập và tạm dừng gói.','Có chức năng và 9 bài thử quy tắc đạt.'),
('Sprint 5','Báo cáo và khu vực hội viên.','Có chức năng; cần đối chiếu số liệu và quyền.'),
('Sprint 6','Chăm sóc và chi phí.','Có chức năng và 4 bài thử quy tắc đạt.'),
('Kiểm tra dùng chung','Kiểm tra mã dữ liệu và tạo ứng dụng.','3 bài thử mã đạt; kiểm tra mã và tạo bản thành công.')],[1.25,2.65,2.5])
p('Cách chia trên liên kết yêu cầu với phần sản phẩm tương ứng, giúp nhận biết việc đã có và việc cần xác minh tiếp. Tổng 16 bài tự động là bằng chứng kỹ thuật hiện có; chưa thể dùng để tính tỷ lệ nghiệm thu của sáu Sprint.')
h('Giá trị đối với môn Công nghệ phần mềm')
p('Việc áp dụng thể hiện ở cách chọn ưu tiên, đặt mục tiêu ngắn hạn, xác định tiêu chí trước khi làm và dùng kết quả kiểm tra để điều chỉnh kế hoạch. Các Sprint tạo ra phần việc từ giao diện đến dữ liệu, giúp phát hiện vấn đề ngay khi chức năng được hoàn thiện.')
p('Để chứng minh quá trình triển khai thực tế, cần lưu thêm bảng việc, phân công, biên bản đánh giá và thay đổi sau mỗi Sprint. Hồ sơ đó bổ sung cho mã nguồn, giúp giải thích không chỉ sản phẩm làm được gì mà còn nhóm đã tổ chức công việc như thế nào.')

# 48–52: conclusion, references and useful appendices.
page('7 Kết luận', True)
h('7.1 Kết quả đạt được')
p('Đề tài đã xây dựng được các phần chính của website quản lý phòng gym một chi nhánh: hội viên, gói tập, bán và gia hạn, khoản thu, vào tập, báo cáo, cổng hội viên, chăm sóc và chi phí. Ba nhóm người dùng được phân biệt theo quyền sử dụng.')
p('Báo cáo trình bày từ bài toán đến yêu cầu, thiết kế, xây dựng và kiểm thử. Các sơ đồ giải thích vai trò của người dùng, dữ liệu cần lưu và cách xử lý những thao tác quan trọng. Bảng kiểm thử liên kết với các trường hợp có thể quan sát được.')
p('Phần Scrum xây dựng danh sách yêu cầu ưu tiên và phương án sáu Sprint, với mục tiêu, kế hoạch, tiêu chí nghiệm thu và kết quả sản phẩm tương ứng. Nội dung giúp thể hiện cách vận dụng quy trình vào đề tài thay vì chỉ nêu tên phương pháp.')
h('7.2 Hạn chế')
p('Kiểm thử tự động mới tập trung vào một số quy tắc; chưa bao phủ toàn bộ thao tác với cơ sở dữ liệu và trình duyệt. Chưa có kết quả đầy đủ về tải, sao lưu và khôi phục. Những giới hạn này cần được xử lý trước khi sử dụng chính thức.')
p('Phạm vi hiện tại là một chi nhánh, thanh toán do nhân viên ghi nhận. Hệ thống chưa tự đối soát ngân hàng hoặc kết nối cửa kiểm soát. Hồ sơ quá trình Scrum còn cần bổ sung bằng chứng thực hiện, phân công và nghiệm thu theo từng Sprint.')
page('7.3 Hướng phát triển')
table(['Thứ tự','Việc cần làm','Lý do ưu tiên'],[
('1','Kiểm tra đầy đủ quyền, bán gói, vào tập và hủy thu.','Giảm nguy cơ sai dữ liệu và vượt quyền.'),
('2','Bổ sung thử tự động cho toàn bộ thao tác chính.','Dễ kiểm tra lại khi sửa chương trình.'),
('3','Hoàn thiện sao lưu, khôi phục và theo dõi lỗi.','Chuẩn bị sử dụng ổn định.'),
('4','Đo tốc độ với dữ liệu lớn hơn và điều chỉnh.','Giữ thao tác tại quầy thuận tiện.'),
('5','Xem xét thanh toán trực tuyến và nhiều chi nhánh.','Mở rộng khi có nhu cầu và quy định rõ.')],[.7,3.6,2.1])
h('Định hướng hoàn thiện quy trình')
p('Ở các đợt tiếp theo, nhóm cần duy trì cùng một cách ghi yêu cầu và bằng chứng kiểm tra. Khi phát sinh thay đổi, cập nhật cả mô tả chức năng, kế hoạch và tình huống kiểm thử để các phần không lệch nhau.')
p('Kết quả cần ưu tiên là một luồng vận hành có thể kiểm chứng từ đầu đến cuối. Sau khi mức chất lượng này ổn định, việc thêm chức năng sẽ có cơ sở rõ hơn và giảm nguy cơ ảnh hưởng đến phần đang dùng.')
page('Tài liệu tham khảo', True)
refs=[
'[1] Ken Schwaber và Jeff Sutherland. The Scrum Guide, 2020. https://scrumguides.org/scrum-guide.html. Dùng cho trách nhiệm, Sprint và các buổi làm việc trong chương 6.',
'[2] Agile Manifesto. Tuyên ngôn Phát triển Phần mềm Linh hoạt. https://agilemanifesto.org/iso/vi/manifesto.html.',
'[3] Next.js. Tài liệu hướng dẫn xây dựng ứng dụng. https://nextjs.org/docs.',
'[4] React. Tài liệu hướng dẫn xây dựng giao diện. https://react.dev/learn.',
'[5] Supabase. Tài liệu đăng nhập và quản lý dữ liệu. https://supabase.com/docs.',
'[6] PostgreSQL. Tài liệu cơ sở dữ liệu. https://www.postgresql.org/docs/.',
'[7] Vitest. Tài liệu kiểm thử. https://vitest.dev/guide/.',
'[8] Dự án GymFlow. README, mã nguồn trong app, components, lib và các tệp dữ liệu trong supabase. Phiên bản tại thời điểm lập báo cáo.',
'[9] Dự án GymFlow. Các tệp kiểm thử lib/check-in.test.ts, lib/operations.test.ts và lib/validation.test.ts; kết quả kiểm chứng được ghi nhận ngày 22 tháng 9 năm 2026.'
]
for ref in refs: p(ref)
page('Phụ lục A Kịch bản trình diễn', True)
p('Chuẩn bị môi trường thử, tài khoản quản lý, nhân viên và hội viên cùng dữ liệu mẫu. Dùng dữ liệu riêng để không làm thay đổi hồ sơ vận hành. Trước buổi trình diễn, xác nhận đăng nhập và kết nối dữ liệu hoạt động.')
table(['Bước','Thao tác','Điều cần quan sát'],[
('1','Đăng nhập bằng quản lý và nhân viên.','Khác biệt về quyền và khu vực sử dụng.'),
('2','Tạo hội viên, tìm và mở hồ sơ.','Thông tin và mã nhận diện được lưu.'),
('3','Bán một gói có giới hạn lượt.','Gói, khoản thu và số lượt ban đầu.'),
('4','Ghi nhận vào tập rồi quét lại.','Chỉ lượt hợp lệ làm giảm số lượt.'),
('5','Thử tạm dừng hoặc gói hết hạn.','Hệ thống giải thích lý do từ chối.'),
('6','Gia hạn và kiểm tra gói kế tiếp.','Ngày hiệu lực phù hợp quy tắc.'),
('7','Mở cổng hội viên.','Chỉ có dữ liệu của tài khoản hiện tại.'),
('8','Xem thu chi và lịch sử thao tác.','Số tổng đối chiếu được với dữ liệu.')],[.6,3,2.8])
p('Nếu một bước không đạt, ghi lại dữ liệu và trạng thái thay vì bỏ qua. Khi báo cáo kết quả, nêu rõ bước đã thử thành công, bước chưa thử và lỗi còn lại. Cách trình diễn này hỗ trợ cả đánh giá sản phẩm và buổi đánh giá cuối Sprint.')
page('Phụ lục B Đối chiếu yêu cầu và sản phẩm', True)
table(['Yêu cầu','Phần sản phẩm','Kiểm tra liên quan'],[
('YC01–02','Đăng nhập, tài khoản và quyền.','Đúng vai trò, khóa tài khoản, đổi mật khẩu.'),
('YC03–04','Hội viên và danh mục gói.','Dữ liệu thiếu, trùng, giá và quyền chỉnh sửa.'),
('YC05–06','Bán mới và gia hạn.','Gói đi cùng khoản thu, ngày hiệu lực.'),
('YC07–08','Vào tập và tạm dừng.','Chín quy tắc tự động và thử dữ liệu thực tế.'),
('YC09','Hủy khoản thu.','Quyền quản lý, lý do và lịch sử.'),
('YC10–11','Báo cáo và khu vực hội viên.','Đối chiếu số tổng, chặn đọc hồ sơ người khác.'),
('YC12–13','Chăm sóc và chi phí.','Bốn bài tự động và thử thao tác đầy đủ.')],[1.05,2.35,3])
h('Hồ sơ cần lưu cho từng Sprint')
bullets(['Mục tiêu và danh sách yêu cầu đã chọn.', 'Nhiệm vụ, người thực hiện và thay đổi kế hoạch.', 'Kết quả kiểm tra gắn với tiêu chí nghiệm thu.', 'Phản hồi khi trình diễn và công việc còn lại.', 'Một hoặc hai hành động cải tiến cho Sprint sau.'])
p('Bảng đối chiếu giúp tìm từ một yêu cầu đến phần sản phẩm và cách kiểm tra. Khi nộp báo cáo, thông tin thành viên trên bìa và hồ sơ thực hiện Sprint cần phản ánh đúng bài làm thực tế.')

assert len(pages)==52, len(pages)
doc.core_properties.title='Xây dựng hệ thống quản lý phòng gym trên nền tảng web theo phương pháp Agile Scrum'
doc.core_properties.subject='Báo cáo Công nghệ phần mềm bản cô đọng tập trung áp dụng Scrum'
doc.core_properties.author='Sinh viên thực hiện'
for t in doc.tables:
    b.set_table_borders(t)
    for row in t.rows:
        for cell in row.cells:
            b.set_cell_margins(cell,top=90,bottom=90)
doc.save(OUT)
(ROOT/'artifacts/report_concise_pages.json').write_text(json.dumps(pages,ensure_ascii=False,indent=2))
print(OUT)
